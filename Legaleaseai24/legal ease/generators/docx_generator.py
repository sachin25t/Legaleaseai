from io import BytesIO

from docx import Document


def format_docx(document: str, document_type: str, terms: str, logo_file=None) -> bytes:
    output = BytesIO()
    doc = Document()

    if logo_file is not None:
        try:
            logo_file.seek(0)
            doc.add_picture(logo_file)
        except Exception:
            pass

    doc.add_heading(document_type or "LegalEase Document", level=0)
    for paragraph in document.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    doc.save(output)
    return output.getvalue()
